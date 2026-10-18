#!/usr/bin/env python3
"""Create the illustrated, standalone Word report. Requires python-docx."""
from pathlib import Path
import argparse
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
args=argparse.ArgumentParser();args.add_argument('--repository-url',required=True);a=args.parse_args()
doc=Document();section=doc.sections[0]
section.top_margin=section.bottom_margin=Cm(1.5);section.left_margin=section.right_margin=Cm(1.7)
section.page_width=Cm(21);section.page_height=Cm(29.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
 style=doc.styles[name];style.font.name='Vazirmatn';style._element.get_or_add_rPr().rFonts.set(qn('w:cs'),'Vazirmatn')
 style.font.size=Pt(10.5 if name=='Normal' else 16 if name=='Heading 1' else 13 if name=='Heading 2' else 27)
 style.font.color.rgb=RGBColor.from_string('23464B')
 style.paragraph_format.space_after=Pt(3)
 style.paragraph_format.line_spacing=1.05

def rtl(p):
 p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 b=OxmlElement('w:bidi');b.set(qn('w:val'),'1');p._p.get_or_add_pPr().append(b)
 return p

def para(text,style=None):
 p=rtl(doc.add_paragraph(text,style));return p

def heading(text): return para(text,'Heading 2')
def table(rows):
 t=doc.add_table(rows=0, cols=2);t.style='Table Grid';t.autofit=False
 t.columns[0].width=Cm(4.0);t.columns[1].width=Cm(13.5)
 direction=OxmlElement('w:bidiVisual');direction.set(qn('w:val'),'1');t._tbl.tblPr.append(direction)
 for row in rows:
  cells=t.add_row().cells
  cells[0].width=Cm(4.0);cells[1].width=Cm(13.5)
  for c,text in zip(cells,row):
   p=rtl(c.paragraphs[0]);r=p.add_run(text);r.font.size=Pt(9.5)
 return t

p=doc.add_paragraph('DUNE DASH','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('گزارش پروژهٔ بازی دایناسور — تمرین اول');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
para('نام و نام خانوادگی: امیرحسین قدردان    |    شمارهٔ دانشجویی: 401521489')
para('درس: طراحی بازی‌های کامپیوتری    |    موتور: Unity 2022.3.62f1 LTS    |    زبان: C#')
para('مخزن کد پروژه در GitHub:')
p=doc.add_paragraph(a.repository_url);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
heading('ایده و ویژگی‌های پیاده‌سازی')
para('Dune Dash یک رانر بی‌پایان دوبعدی با تم بیابانی است. بازیکن با پرش و Dash از کاکتوس‌ها و پرندگان عبور می‌کند و سکه جمع می‌کند. سرعت بازی با مسافت افزایش می‌یابد. طراحی پیکسلی، موسیقی و افکت‌ها برای همین پروژه تولید شده‌اند.')
table([
 ('نیاز تمرین','پیاده‌سازی در پروژه'),
 ('پس‌زمینهٔ پویا','Parallax با ابرها، سه لایهٔ تپه، گیاه و زمین متحرک.'),
 ('ورودی جدید یونیتی','InputAction و InputSystemUIInputModule؛ کنترل قدیمی استفاده نشده است.'),
 ('جان، پرش و Dash','۳ جان و نوار سلامت؛ Space برای پرش، Shift برای Dash با شارژ ۲٫۲ ثانیه.'),
 ('برخورد و آیتم','موانع زمینی و هوایی و سکه؛ هر برخورد یک جان و هر سکه ۲۵ امتیاز.'),
 ('طراحی شی‌گرا','GroundObstacle، FlyingObstacle و Coin از کلاس مشترک Interactable ارث می‌برند.'),
 ('Animator','چهار حالت Run، Jump، Hit و Death با Transition؛ تغییر رنگ و چشمک هنگام آسیب.'),
 ('رابط و ذخیره‌سازی','Main Menu و Pause Menu مستقل؛ امتیاز جاری و رکورد زنده؛ ذخیرهٔ رکورد با PlayerPrefs.')])
heading('ساختار و آزمون')
para('GameManager چرخهٔ بازی، امتیاز و رکورد را مدیریت می‌کند. DinoController مسئول حرکت و سلامت است. Interactable قرارداد مشترک تعامل و حرکت را تعریف می‌کند. ObstacleSpawner الگوهای مانع و سکه را تولید می‌کند. GameUI، AudioDirector و VisualEffects رابط، صدا و جلوه‌ها را کنترل می‌کنند.')
para('Build مستقل لینوکس روی همین سیستم اجرا شده و هر ۲۵ آزمون موفق بوده است: ورودی جدید، پرش و جلوگیری از پرش دوم، Dash، سکه، برخورد هر دو نوع مانع، توقف، مرگ، ذخیرهٔ رکورد، شروع مجدد و پیوستگی زمین. نتایج کامل در Documentation/test-results.txt است.')
heading('اجرا و منابع')
para('اجرای Build: فایل RunGame.sh یا Builds/Linux/DuneDash.x86_64. اجرای پروژه در Unity: صحنهٔ Assets/DuneDash/Scenes/Desert.unity و سپس Play. پوشه‌های سنگین Library، Builds و Logs طبق gitignore وارد مخزن نمی‌شوند؛ نسخهٔ اجرایی در ZIP تحویل قرار دارد.')
para('منابع: صورت تمرین HM1.pdf، مستندات رسمی Unity Input System، Animator و PlayerPrefs. فونت Lato با مجوز OFL؛ تصاویر و صداهای اختصاصی. کد و مستندات با کمک OpenAI Codex تهیه شده‌اند؛ انتساب و لینک منابع در Documentation/SOURCES.md درج شده است.')

for title,figs in [
 ('تصاویر واقعی بازی — شروع و گیم‌پلی',[('01-main-menu.png','تصویر ۱ — منوی اصلی و راهنمای کنترل‌ها'),('07-gameplay.png','تصویر ۲ — گیم‌پلی، جان‌ها، امتیاز، موانع و سکه‌ها')]),
 ('تصاویر واقعی بازی — توقف و پایان',[('04-pause.png','تصویر ۳ — منوی توقف با ادامه، شروع مجدد و برگشت'),('06-game-over.png','تصویر ۴ — پایان بازی، امتیاز نهایی و رکورد')])]:
 doc.add_page_break();para(title,'Heading 1')
 for filename,caption in figs:
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.add_run().add_picture(str(root/'Documentation/Screenshots'/filename),width=Cm(17.3))
  p=para(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.space_after=Pt(15)

footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Dune Dash  |  401521489  |  ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.author='امیرحسین قدردان'
doc.core_properties.title='گزارش پروژهٔ Dune Dash — 401521489'
out=root/'Documentation/GD_Hw1_Amirhossein_Ghadrdan_401521489.docx';doc.save(out);print(out)
